from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document

from rag_research_agent.models import SourceDocument
from rag_research_agent.parsing import DocxParsingError, UnsupportedFileTypeError, extract_text

DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_document(content, mime_type=DOCX_MIME_TYPE):
    return SourceDocument(
        id="doc-1",
        source="local",
        title="example",
        uri="/tmp/example.docx",
        content=content,
        mime_type=mime_type,
        last_modified=datetime(2026, 1, 1, tzinfo=UTC),
    )


def to_bytes(word_document):
    buffer = BytesIO()
    word_document.save(buffer)
    return buffer.getvalue()


def test_paragraphs_in_order_separated_by_blank_line():
    word = Document()
    word.add_paragraph("First paragraph")
    word.add_paragraph("Second paragraph")

    assert extract_text(make_document(to_bytes(word))) == "First paragraph\n\nSecond paragraph"


def test_table_row_cells_joined_by_pipe():
    word = Document()
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    table.cell(1, 0).text = "c"
    table.cell(1, 1).text = "d"

    assert extract_text(make_document(to_bytes(word))) == "a | b\nc | d"


def test_paragraph_table_paragraph_order_is_kept():
    word = Document()
    word.add_paragraph("Before")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    word.add_paragraph("After")

    assert extract_text(make_document(to_bytes(word))) == "Before\n\na | b\n\nAfter"


def test_empty_paragraphs_are_skipped():
    word = Document()
    word.add_paragraph("One")
    word.add_paragraph("")
    word.add_paragraph("   ")
    word.add_paragraph("Two")

    assert extract_text(make_document(to_bytes(word))) == "One\n\nTwo"


def test_broken_bytes_raise_clear_error():
    with pytest.raises(DocxParsingError, match="not a valid .docx"):
        extract_text(make_document(b"this is not a docx"))


def test_old_doc_format_is_unsupported():
    with pytest.raises(UnsupportedFileTypeError, match="application/msword"):
        extract_text(make_document(b"\xd0\xcf\x11\xe0", "application/msword"))
